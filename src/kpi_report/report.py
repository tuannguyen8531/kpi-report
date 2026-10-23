"""Report output generation: CSV and Excel files."""

import os
import re
from calendar import monthrange
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import cast

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from .constants import OUTPUT_COLUMNS
from .processor import ReportStats


def format_excel_filename(pattern: str | None, month: int, year: int) -> str:
    """Apply the extension's month/year placeholders to a safe download filename."""
    name = (pattern or '').strip() or 'report_MM_YYYY.xlsx'
    name = name.replace('YYYY', str(year)).replace('MM', f'{month:02d}')
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', name).rstrip(' .')
    return name if name.lower().endswith('.xlsx') else name + '.xlsx'


def save_report(
    output_dfs: list[pd.DataFrame],
    output_file: str,
    stats: ReportStats,
    *,
    month: int,
    year: int,
    off_df: pd.DataFrame | None = None,
    excel_file: str | None = None,
) -> None:
    """Save the final report to CSV and Excel files.

    Args:
        output_dfs: List of project DataFrames (with headers and separators).
        output_file: Path to the output CSV file.
        stats: Accumulated report statistics.
        month: Reporting month for the Excel title.
        year: Reporting year for the Excel title.
        off_df: Optional time-off entries, summarized in the scoring explanation.
        excel_file: Optional Excel output path, independent of the CSV filename.
    """
    os.makedirs(os.path.dirname(output_file), exist_ok=True)
    print("Combining data and creating final report...")
    project_dfs = output_dfs.copy()
    output_dfs = output_dfs.copy()

    # Blank separator before summary
    blank_df = pd.DataFrame([[""] * len(OUTPUT_COLUMNS)], columns=OUTPUT_COLUMNS)
    output_dfs.append(blank_df)

    # Summary rows
    summary_rows = [
        {"Url": "Number of tasks: ", "Start date": str(stats.total_tasks)},
        {"Url": "Total spent time:", "Start date": str(stats.total_spent)},
        {"Url": "Total estimate time:", "Start date": str(stats.total_estimate)},
    ]
    for row in summary_rows:
        for col in OUTPUT_COLUMNS:
            row.setdefault(col, "")
    output_dfs.append(pd.DataFrame(summary_rows))

    final_df = pd.concat(output_dfs, ignore_index=True)

    # Save CSV
    print(f"Saving CSV file: {output_file}")
    final_df.to_csv(output_file, index=False, encoding="utf-8-sig")

    # Save Excel
    excel_file = excel_file or str(Path(output_file).with_suffix('.xlsx'))
    print(f"Creating Excel file: {excel_file}")
    _save_excel(project_dfs, excel_file, month, year, off_df)


def _save_excel(
    project_dfs: list[pd.DataFrame],
    excel_file: str,
    month: int,
    year: int,
    off_df: pd.DataFrame | None = None,
) -> None:
    """Fill the work-report template with native Excel values and KPI formulas."""
    template = Path(__file__).with_name("templates") / "work_report.xlsx"
    workbook = load_workbook(template)
    sheet = cast(Worksheet, workbook["Báo cáo công việc"])
    last_day = monthrange(year, month)[1]
    sheet["B2"] = f"Tháng 01/{month:02d}/{year} - {last_day:02d}/{month:02d}/{year}"

    styles = {
        row: [copy(sheet.cell(row, col)._style) for col in range(2, 11)]
        for row in (4, 5, 6)
    }
    headers = [sheet.cell(5, col).value for col in range(2, 11)]
    for row in sheet.iter_rows(min_row=4, max_row=6, min_col=2, max_col=10):
        for cell in row:
            cell.value = None
            cell._style = None

    next_row = 4

    def write_row(values: list, style_row: int, hyperlink: bool = True) -> None:
        nonlocal next_row
        for col, (value, style) in enumerate(zip(values, styles[style_row]), start=2):
            cell = cast(Cell, sheet.cell(next_row, col))
            cell._style = copy(style)
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"
            if style_row == 6:
                if col in (3, 4, 5):
                    cell.number_format = "dd/mm/yyyy"
                if (
                    col == 2
                    and hyperlink
                    and isinstance(value, str)
                    and value.startswith(("https://", "http://"))
                ):
                    cell.hyperlink = value
        sheet.row_dimensions[next_row].height = sheet.row_dimensions[style_row].height
        next_row += 1

    def write_data(frame: pd.DataFrame) -> None:
        for record in frame[OUTPUT_COLUMNS].itertuples(index=False, name=None):
            values = []
            for index, value in enumerate(record):
                if pd.isna(value) or value == "":
                    value = None
                elif index in (1, 2, 3):
                    date_month, date_day, date_year = map(int, str(value).split("/"))
                    value = date(date_year, date_month, date_day)
                elif index in (4, 5, 6):
                    value = float(value)
                values.append(value)
            write_row(values, 6, hyperlink=not frame.attrs.get('manual', False))

    for frame in project_dfs:
        if frame.empty:
            continue
        # The processor emits a project title, its tasks, then a blank separator.
        if (frame.loc[:, OUTPUT_COLUMNS[1:]].to_numpy() == "").all():
            title = frame.iloc[0]["Url"]
            if title:
                write_row([title] + [None] * 8, 4)
                write_row(headers, 5)
            else:
                next_row += 1
        else:
            write_data(frame)

    # Formula ranges grow with the exported project rows.
    last_task_row = max(next_row - 1, 6)
    task_types = f"I6:I{last_task_row}"
    for cell_row, task_type in ((6, "Kế hoạch"), (7, "Phát sinh")):
        sheet.cell(cell_row, 13, f'=COUNTIF({task_types},"{task_type}")')
    sheet["M5"] = "=M6+M7"
    sheet["M9"] = f"=SUM(F6:F{last_task_row})"
    sheet["M10"] = f"=SUM(G6:G{last_task_row})"
    sheet["M11"] = "=M10-M12"
    sheet["M12"] = f'=SUMIF({task_types},"Phát sinh",G6:G{last_task_row})'
    for cell_row, column, criterion in (
        (13, "C", ""),
        (14, "D", ""),
        (15, "F", ""),
        (16, "G", ""),
        (17, "J", "Đúng hạn"),
        (18, "J", "Trễ hạn"),
        (19, "H", "0"),
        (20, "H", ">0"),
    ):
        value_range = f"{column}6:{column}{last_task_row}"
        counts = [
            f'COUNTIFS({task_types},"{task_type}",{value_range},"{criterion}")'
            for task_type in ("Kế hoạch", "Phát sinh")
        ]
        sheet.cell(cell_row, 13, "=" + "+".join(counts))

    for row in range(24, 37):
        cell = cast(Cell, sheet.cell(row, 13))
        formula = cell.value
        if not isinstance(formula, str) or not formula.startswith("="):
            raise ValueError(f"Missing formula in template cell {cell.coordinate}")
        cell.value = f"=IFERROR({formula[1:]},0)"

    # Leave belongs only in the scoring explanation, not in the work task table.
    if off_df is not None and not off_df.empty:
        reasons = {}
        for row in off_df.itertuples(index=False, name=None):
            reason, start, _, _, _, spent, *_ = row
            group = reasons.setdefault(reason or 'OFF', {'hours': 0, 'dates': set()})
            group['hours'] += float(spent)
            group['dates'].add(datetime.strptime(start, '%m/%d/%Y').strftime('%d/%m'))
        lines = []
        for index, (reason, group) in enumerate(reasons.items(), start=1):
            hours = f"{group['hours']:.8f}".rstrip('0').rstrip('.')
            lines.append(f"{index}. {reason} {hours}h ({', '.join(sorted(group['dates']))})")
        scoring = cast(Worksheet, workbook['Chấm điểm KPI'])
        cell = scoring['A53']
        cell.value = '\n'.join(lines)
        cell.data_type = 's'
        alignment = copy(cell.alignment)
        alignment.wrapText = True
        alignment.vertical = 'top'
        cell.alignment = alignment
        text_lines = sum(max(1, (len(line) + 99) // 100) for line in cell.value.split('\n'))
        scoring.row_dimensions[53].height = max(scoring.row_dimensions[53].height or 0, text_lines * 18 + 12)

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(excel_file)
