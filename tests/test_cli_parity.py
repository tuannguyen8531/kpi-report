"""The CLI and extension must produce the same workbook from the same inputs."""

import json
import subprocess
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

import pandas as pd
from openpyxl import load_workbook

from kpi_report.cli import parse_arguments
from kpi_report.config import validate_projects_in_csv, load_url_map
from kpi_report.processor import (
    create_off_dataframe, create_project_dataframes, process_off_data, process_project_data,
)
from kpi_report.report import format_excel_filename, save_report
from kpi_report.main import main, load_and_validate_data

ROOT = Path(__file__).resolve().parents[1]


class CLIParityTest(unittest.TestCase):
    def test_workbook_matches_extension(self):
        self.assertEqual(
            (ROOT / 'src/kpi_report/templates/work_report.xlsx').read_bytes(),
            (ROOT / 'extension/assets/templates/work_report.xlsx').read_bytes(),
        )
        raw = pd.DataFrame([
            ['Shared (team-a/app)', '12', 'TASK', 2, '2026-10-01'],
            ['Shared (team-b/app)', '187', 'MR', 1, '2026-10-02'],
            ['Nghiên cứu', 'https://manual.example.com', 'MANUAL', 1.5, '2026-10-03'],
            ['Nghiên cứu', 'https://manual.example.com', 'MANUAL', 2, '2026-10-04'],
            ['OFF', 'Nghỉ lễ', 'OFF', 8, '2026-10-05'],
            ['OFF', 'Nghỉ lễ', 'OFF', 4, '2026-10-06'],
        ], columns=['Project', 'Task', 'Type', 'Time', 'Date'])
        logs = [dict(project=r.Project, task=r.Task, title=r.Task if r.Type == 'MANUAL' else '',
                     type=r.Type, hours=r.Time, date=r.Date) for r in raw.itertuples() if r.Type != 'OFF']
        leaves = [dict(day=r.Date, reason=r.Task, hours=r.Time) for r in raw.itertuples() if r.Type == 'OFF']
        raw['Date'] = pd.to_datetime(raw['Date'])
        url_map = {}
        validate_projects_in_csv(raw, url_map)
        self.assertEqual(url_map, {'Shared (team-a/app)': 'team-a/app', 'Shared (team-b/app)': 'team-b/app'})
        client = Mock()
        client.get_work_item.return_value = {'estimate': '2.00', 'task_type': 'Phát sinh'}
        client.get_merge_request.return_value = {'estimate': '1.00'}
        frames, stats = create_project_dataframes(process_project_data(raw), url_map, client, 'https://gitlab.example.com')
        self.assertEqual(stats.total_tasks, 3)
        self.assertEqual(stats.total_spent, 6.5)
        self.assertEqual(client.get_work_item.call_count, 1)
        self.assertEqual(client.get_merge_request.call_count, 1)
        self.assertEqual(frames[-3].iloc[0]['Url'], 'Nghiên cứu')
        with TemporaryDirectory() as directory:
            cli = Path(directory) / 'cli.csv'
            extension = Path(directory) / 'extension.xlsx'
            save_report(frames, str(cli), stats, month=10, year=2026, off_df=create_off_dataframe(process_off_data(raw)))
            script = r"""
import fs from 'node:fs';
import './extension/src/vendor/exceljs.min.js';
import {processProjectData, processOffData, buildProjectData, generateExcelWorkbook} from './extension/src/excel.js';
const input = JSON.parse(fs.readFileSync(0, 'utf8'));
const grouped = processProjectData(input.logs);
const enrichment = new Map(grouped.filter(r => r.Type !== 'MANUAL').map(r => [
  `${r.Project}:::${r.Task}:::${r.Type}`, r.Type === 'TASK'
    ? {estimate: '2.00', task_type: 'Phát sinh'} : {estimate: '1.00'},
]));
const {projectBlocks} = buildProjectData(grouped, enrichment, {url: 'https://gitlab.example.com', projects: input.projects});
const result = await generateExcelWorkbook(fs.readFileSync('extension/assets/templates/work_report.xlsx'), {
  month: 10, year: 2026, projectBlocks, offEntries: processOffData(input.leaves),
}, globalThis.ExcelJS);
fs.writeFileSync(input.output, Buffer.from(result));
"""
            subprocess.run(['node', '--input-type=module', '-e', script], cwd=ROOT, check=True,
                           input=json.dumps({'logs': logs, 'leaves': leaves, 'projects': [
                               {'project': name, 'url': path} for name, path in url_map.items()
                           ], 'output': str(extension)}), text=True, capture_output=True)
            python_book = load_workbook(cli.with_suffix('.xlsx'))
            js_book = load_workbook(extension)
            try:
                for sheet in python_book:
                    other = js_book[sheet.title]
                    for row in sheet:
                        for cell in row:
                            self.assertEqual(cell.value, other[cell.coordinate].value, f'{sheet.title}!{cell.coordinate}')
                            self.assertEqual(cell.hyperlink.target if cell.hyperlink else None,
                                             other[cell.coordinate].hyperlink.target if other[cell.coordinate].hyperlink else None)
                self.assertIsNone(python_book['Báo cáo công việc']['B14'].hyperlink)
                self.assertEqual(python_book['Chấm điểm KPI']['A53'].value, '1. Nghỉ lễ 12h (05/10, 06/10)')
            finally:
                python_book.close()
                js_book.close()

    def test_filename_and_arguments(self):
        self.assertEqual(format_excel_filename('Report - KPI (MM/YYYY)', 10, 2026), 'Report - KPI (10_2026).xlsx')
        self.assertEqual(format_excel_filename(None, 2, 2024), 'report_02_2024.xlsx')
        with patch('sys.argv', ['report', '-m', '10', '-y', '2026', '--excel-pattern', 'Work (MM-YYYY)']):
            args = parse_arguments()
        self.assertEqual((args.month, args.year, args.excel_pattern), (10, 2026, 'Work (MM-YYYY)'))

    def test_manual_only_cli_needs_no_gitlab_and_filters_reporting_month(self):
        with TemporaryDirectory() as directory:
            source = Path(directory) / 'tasks.csv'
            output = Path(directory) / 'report.csv'
            source.write_text('Project,Task,Type,Time,Date\n'
                              'Công việc ngoài GitLab,Research,MANUAL,2,"October 1, 2026"\n'
                              'Nghiên cứu,Research,MANUAL,3,"October 2, 2026"\n'
                              'Nghiên cứu,Other month,MANUAL,8,"November 1, 2026"\n'
                              'OFF,Nghỉ phép,OFF,4,"October 3, 2026"\n', encoding='utf-8')
            with patch('sys.argv', ['report', '-m', '10', '-y', '2026', '--excel-pattern', 'Report (MM/YYYY)']), \
                    patch('kpi_report.main.load_env'), patch('kpi_report.main.load_url_map', return_value={}), \
                    patch('kpi_report.main.get_file_paths', return_value=(str(source), str(output))), \
                    patch.dict('os.environ', {}, clear=True):
                main()
            book = load_workbook(Path(directory) / 'Report (10_2026).xlsx')
            try:
                self.assertEqual(book['Báo cáo công việc']['B4'].value, 'Nghiên cứu')
                self.assertEqual(book['Báo cáo công việc']['B6'].value, 'Research')
                self.assertEqual(book['Báo cáo công việc']['G6'].value, 5)
                self.assertEqual(book['Chấm điểm KPI']['A53'].value, '1. Nghỉ phép 4h (03/10)')
                self.assertFalse(any(c.value == 'Other month' for c in book['Báo cáo công việc']['B']))
            finally:
                book.close()

    def test_invalid_csv_and_optional_project_config(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            with patch('kpi_report.config.get_project_root', return_value=root):
                self.assertEqual(load_url_map(), {})
            source = root / 'input.csv'
            for kind, title, hours in [('MANUAL', 'Work', '0'), ('MANUAL', 'Work', '25'),
                                       ('OFF', 'Leave', 'inf'), ('UNKNOWN', 'Work', '1'),
                                       ('MANUAL', '   ', '2')]:
                with self.subTest(kind=kind, title=title, hours=hours):
                    source.write_text(f'Project,Task,Type,Time,Date\nNghiên cứu,{title},{kind},{hours},"October 1, 2026"\n')
                    with self.assertRaises(ValueError):
                        load_and_validate_data(str(source), {})
