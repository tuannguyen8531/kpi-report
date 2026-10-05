"""Run with: python -m unittest discover -s tests."""

import unittest
from datetime import date, datetime
from itertools import product
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import cast

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from kpi_report.constants import OUTPUT_COLUMNS
from kpi_report.processor import (
    ReportStats,
    create_off_dataframe,
    process_off_data,
    process_project_data,
)
from kpi_report.report import save_report


class ReportTest(unittest.TestCase):
    def test_template_export(self):
        blank = pd.DataFrame([[""] * 9], columns=OUTPUT_COLUMNS)
        header = blank.copy()
        header.at[0, "Url"] = "Project A"
        task = [
            "https://gitlab.example.com/a/-/work_items/1",
            "02/01/2024",
            "",
            "02/03/2024",
            "4.50",
            3.25,
            0,
            "Kế hoạch",
            "Đúng hạn",
        ]
        tasks = pd.DataFrame([task], columns=OUTPUT_COLUMNS)
        other_header = header.copy()
        other_header.at[0, "Url"] = "Project B"
        other = tasks.copy()
        other.loc[0, ["Url", "Estimate", "Reopen count", "Task Type", "Progress"]] = [
            "",
            "",
            2,
            "Phát sinh",
            "Trễ hạn",
        ]
        raw = pd.DataFrame(
            [
                ["Project A", "1", "TASK", 3.25, "2024-02-01"],
                ["OFF", "Vacation", "OFF", 8, "2024-02-04"],
            ],
            columns=["Project", "Task", "Type", "Time", "Date"],
        )
        raw["Date"] = pd.to_datetime(raw["Date"])
        self.assertEqual(process_project_data(raw)["Project"].tolist(), ["Project A"])
        off_df = create_off_dataframe(process_off_data(raw))
        template_workbook = load_workbook(
            Path(__file__).resolve().parents[1]
            / "src/kpi_report/templates/work_report.xlsx"
        )
        template = cast(Worksheet, template_workbook["Báo cáo công việc"])

        with TemporaryDirectory() as directory:
            for count, include_off in product((0, 1, 1005), (False, True)):
                with self.subTest(task_count=count, include_off=include_off):
                    frames = (
                        []
                        if not count
                        else [
                            header,
                            pd.concat([tasks] * count, ignore_index=True),
                            blank,
                            other_header,
                            other,
                            blank,
                        ]
                    )
                    path = Path(directory) / f"report_{count}.csv"
                    save_report(
                        frames,
                        str(path),
                        ReportStats(count + bool(count), count * 3.25, count * 4.5),
                        month=2,
                        year=2024,
                        off_df=off_df if include_off else None,
                    )
                    # Exporting does not append summary rows to the caller's list.
                    self.assertEqual(len(frames), 6 if count else 0)
                    workbook = load_workbook(path.with_suffix(".xlsx"))
                    self.assertEqual(workbook.sheetnames, ["Báo cáo công việc"])
                    sheet = cast(Worksheet, workbook["Báo cáo công việc"])
                    self.assertEqual(sheet["B2"].value, "Tháng 01/02/2024 - 29/02/2024")
                    self.assertEqual(sheet["M8"].value, 192)
                    self.assertEqual(
                        str(sheet.merged_cells), str(template.merged_cells)
                    )
                    for col in ("B", "L", "M"):
                        self.assertEqual(
                            sheet.column_dimensions[col].width,
                            template.column_dimensions[col].width,
                        )
                    end = count + 10 if count else 6
                    self.assertEqual(sheet["M10"].value, f"=SUM(G6:G{end})")
                    self.assertEqual(sheet["M9"].value, f"=SUM(F6:F{end})")
                    self.assertEqual(sheet["M5"].value, "=M6+M7")
                    self.assertEqual(
                        sheet["M7"].value, f'=COUNTIF(I6:I{end},"Phát sinh")'
                    )
                    self.assertEqual(
                        sheet["M12"].value, f'=SUMIF(I6:I{end},"Phát sinh",G6:G{end})'
                    )
                    self.assertEqual(
                        sheet["M14"].value,
                        f'=COUNTIFS(I6:I{end},"Kế hoạch",D6:D{end},"")'
                        f'+COUNTIFS(I6:I{end},"Phát sinh",D6:D{end},"")',
                    )
                    self.assertEqual(
                        sheet["M21"].value,
                        f'=COUNTIFS(I6:I{end},"Kế hoạch",H6:H{end},">0")'
                        f'+COUNTIFS(I6:I{end},"Phát sinh",H6:H{end},">0")',
                    )
                    for row in (25, 26, 27, 28, 29, 31, 32, 33, 34, 36, 37, 38, 39):
                        formula = template.cell(row, 13).value
                        assert isinstance(formula, str)
                        self.assertEqual(
                            sheet.cell(row, 13).value, f"=IFERROR({formula[1:]},0)"
                        )
                    if include_off:
                        off_row = next(c.row for c in sheet["B"] if c.value == "OFF")
                        self.assertGreater(off_row, end)
                        self.assertEqual(sheet.cell(off_row + 1, 2).value, "Tasks ")
                        self.assertEqual(sheet.cell(off_row + 2, 2).value, "Vacation")
                        self.assertEqual(sheet.cell(off_row + 2, 7).value, 8)
                        off_date = sheet.cell(off_row + 2, 3).value
                        assert isinstance(off_date, datetime)
                        self.assertEqual(off_date.date(), date(2024, 2, 4))
                    else:
                        self.assertFalse(
                            any(c.value in ("OFF", "Vacation") for c in sheet["B"])
                        )
                    self.assertEqual(workbook.calculation.calcMode, "auto")
                    if count:
                        start_date = sheet["C6"].value
                        assert isinstance(start_date, datetime)
                        self.assertEqual(start_date.date(), date(2024, 2, 1))
                        self.assertEqual(sheet["C6"].number_format, "dd/mm/yyyy")
                        self.assertIsNone(sheet["D6"].value)
                        self.assertEqual(sheet["F6"].value, 4.5)
                        self.assertEqual(sheet["G6"].value, 3.25)
                        hyperlink = sheet["B6"].hyperlink
                        assert hyperlink is not None
                        self.assertEqual(hyperlink.target, task[0])
                        for coordinate in ("B4", "B5", "B6", "L4", "M25"):
                            self.assertEqual(
                                sheet[coordinate]._style, template[coordinate]._style
                            )
                    # CSV retains its original columns and summary section.
                    csv = pd.read_csv(path)
                    self.assertEqual(list(csv.columns), OUTPUT_COLUMNS)
                    self.assertEqual(csv.iloc[-3]["Url"], "Number of tasks: ")
                    self.assertEqual(float(csv.iloc[-2]["Start date"]), count * 3.25)
                    self.assertFalse(csv["Url"].isin(["OFF", "Vacation"]).any())
                    workbook.close()
        template_workbook.close()


if __name__ == "__main__":
    unittest.main()
